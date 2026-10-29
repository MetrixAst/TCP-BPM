import io

import openpyxl
from django.core.cache import cache
from django.test import TestCase
from django.urls import reverse

from account.models import UserAccount, Department
from hr.models import Company
from tickets.models import ServiceRequest, ServiceRequestHistory


class TicketsReportTest(TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        company = Company.objects.create(name='Rep Co', bin_number='343434343434')
        self.ops = Department.objects.create(name='Эксплуатация', company=company)
        self.owner = UserAccount.objects.create_user(username='owner_rep', password='pass', role='owner')
        self.staff = UserAccount.objects.create_user(username='staff_rep', password='pass', role='staff')

        def make(status, category='electrical', department=None):
            return ServiceRequest.objects.create(
                title=f'{status} {category}', description='-', category=category,
                status=status, department=department,
            )

        make('new')
        make('in_progress', department=self.ops)
        make('pending_approval', department=self.ops)
        done = make('done', category='plumbing', department=self.ops)
        ServiceRequestHistory.objects.create(request=done, status='done')
        make('rejected')
        make('cancelled', category='plumbing')

    def _groups(self, response):
        return {g['key']: g['count'] for g in response.context['groups']}

    def test_status_groups(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('tickets:report'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._groups(response), {'created': 1, 'in_work': 2, 'done': 1, 'rejected': 2})
        self.assertEqual(response.context['total'], 6)

    def test_filters(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('tickets:report'), {'department': self.ops.pk, 'category': 'electrical'})
        self.assertEqual(self._groups(response), {'created': 0, 'in_work': 2, 'done': 0, 'rejected': 0})

        response = self.client.get(reverse('tickets:report'), {'group': 'rejected'})
        self.assertEqual(len(response.context['rows']), 2)

    def test_excel_export(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('tickets:report_export'))
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Сводка', 'Заявки'])
        self.assertEqual(wb['Заявки'].max_row, 7)

    def test_staff_without_reports_forbidden(self):
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse('tickets:report')).status_code, 403)
