"""
Массовый импорт справочников из Excel.

Файл проверяется целиком: если хотя бы в одной строке ошибка, ничего не
сохраняется — пользователь видит список ошибок по номерам строк.
Существующие записи обновляются (арендатор — по названию, контрагент — по БИН/ИИН).
"""

from datetime import date, datetime
from io import BytesIO

import openpyxl
from django.db import transaction
from openpyxl.styles import Font


class ImportResult:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.errors = []  # (номер строки, текст)

    @property
    def ok(self):
        return not self.errors


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value, row, label, errors, required=False):
    if value in (None, ''):
        if required:
            errors.append((row, f'«{label}» обязательно'))
        return None
    try:
        return float(str(value).replace(' ', '').replace(',', '.'))
    except ValueError:
        errors.append((row, f'«{label}»: ожидается число, получено «{value}»'))
        return None


def _date(value, row, label, errors):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    errors.append((row, f'«{label}»: дата в формате дд.мм.гггг, получено «{value}»'))
    return None


def _rows(file, columns):
    """Строки листа как dict по заголовкам; заголовки сопоставляются без учёта регистра."""
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [(_text(h)).lower() for h in next(rows, [])]
    index = {}
    for key, title in columns:
        if title.lower() in header:
            index[key] = header.index(title.lower())
    for number, values in enumerate(rows, start=2):
        if not any(v not in (None, '') for v in values):
            continue
        yield number, {key: values[i] if i < len(values) else None for key, i in index.items()}


def _template(columns, example):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([title for _key, title in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(example)
    for i in range(len(columns)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = 22
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ─────────────────────────── Арендаторы ───────────────────────────

TENANT_COLUMNS = [
    ('name', 'Название'),
    ('room', 'Помещение'),
    ('category', 'Категория'),
    ('area', 'Площадь'),
    ('price', 'Ставка за м2'),
    ('discount', 'Скидка %'),
    ('percent', 'Процент с оборота'),
    ('start_date', 'Начало аренды'),
    ('end_date', 'Окончание аренды'),
    ('contact', 'Ответственное лицо'),
    ('phone', 'Телефон'),
    ('email', 'Email'),
    ('address', 'Адрес'),
    ('note', 'Примечание'),
]


def tenant_template():
    return _template(TENANT_COLUMNS, [
        'ТОО «Кофейня»', 'A-101', 'Кафе', 45, 6500, 0, 0,
        '01.01.2026', '31.12.2026', 'Иванов И.', '+7 700 000 00 00', 'cafe@example.kz', 'г. Алматы', '',
    ])


def import_tenants(file):
    from tenants.models import Room, Tenant, TenantCategory

    result = ImportResult()
    prepared = []
    for row, data in _rows(file, TENANT_COLUMNS):
        errors = []
        name = _text(data.get('name'))
        if not name:
            errors.append((row, '«Название» обязательно'))
        values = {
            'area': _number(data.get('area'), row, 'Площадь', errors, required=True),
            'price': _number(data.get('price'), row, 'Ставка за м2', errors, required=True),
            'discount': int(_number(data.get('discount'), row, 'Скидка %', errors) or 0),
            'percent': int(_number(data.get('percent'), row, 'Процент с оборота', errors) or 0),
            'start_date': _date(data.get('start_date'), row, 'Начало аренды', errors),
            'end_date': _date(data.get('end_date'), row, 'Окончание аренды', errors),
            'contact': _text(data.get('contact'))[:40],
            'phone': _text(data.get('phone'))[:40],
            'email': _text(data.get('email'))[:40],
            'address': _text(data.get('address'))[:100],
            'note': _text(data.get('note'))[:300],
        }
        if values['email'] and '@' not in values['email']:
            errors.append((row, f'«Email»: некорректный адрес «{values["email"]}»'))
        result.errors.extend(errors)
        prepared.append((name, _text(data.get('room')), _text(data.get('category')) or 'Прочее', values))

    if not result.ok:
        return result

    with transaction.atomic():
        for name, room_number, category_title, values in prepared:
            if room_number:
                values['room'], _ = Room.objects.get_or_create(
                    number=room_number, defaults={'map_id': room_number, 'floor': 0},
                )
            category = TenantCategory.objects.filter(title__iexact=category_title).first()
            values['category'] = category or TenantCategory.objects.create(title=category_title)

            tenant = Tenant.objects.filter(name__iexact=name).first()
            if tenant is None:
                Tenant.objects.create(name=name, **values)
                result.created += 1
            else:
                for field, value in values.items():
                    setattr(tenant, field, value)
                tenant.save()
                result.updated += 1
    return result


# ─────────────────────────── Контрагенты ───────────────────────────

SUPPLIER_COLUMNS = [
    ('name', 'Наименование'),
    ('identifier', 'БИН/ИИН'),
    ('supplier_type', 'Юр./физ. лицо'),
    ('address1', 'Юридический адрес'),
    ('address2', 'Фактический адрес'),
    ('head_name', 'Руководитель'),
    ('phone', 'Телефон'),
    ('email', 'Email'),
    ('contacts', 'Контакты'),
    ('category', 'Категория'),
]


def supplier_template():
    return _template(SUPPLIER_COLUMNS, [
        'ТОО «Ромашка»', '123456789012', 'Юр. лицо', 'г. Алматы, ул. Абая 1', 'г. Алматы, ул. Абая 1',
        'Петров П.', '+7 700 000 00 00', 'info@example.kz', 'Менеджер Анна', 'Клининг',
    ])


def import_suppliers(file, author=None):
    from purchases.enums import SupplierStatusEnum, SupplierTypeEnum
    from purchases.models import Supplier

    result = ImportResult()
    prepared = []
    for row, data in _rows(file, SUPPLIER_COLUMNS):
        errors = []
        name = _text(data.get('name'))[:100]
        identifier = _text(data.get('identifier'))
        if not name:
            errors.append((row, '«Наименование» обязательно'))
        if identifier and (not identifier.isdigit() or len(identifier) != 12):
            errors.append((row, f'«БИН/ИИН»: 12 цифр, получено «{identifier}»'))
        kind = _text(data.get('supplier_type')).lower()
        values = {
            'name': name,
            'supplier_type': (
                SupplierTypeEnum.INDIVIDUAL.value[0] if kind.startswith('физ') else SupplierTypeEnum.LEGAL.value[0]
            ),
            'address1': _text(data.get('address1'))[:60] or None,
            'address2': _text(data.get('address2'))[:60] or _text(data.get('address1'))[:60] or '—',
            'head_name': _text(data.get('head_name'))[:40] or None,
            'phone': _text(data.get('phone'))[:40] or None,
            'email': _text(data.get('email'))[:40] or None,
            'contacts': _text(data.get('contacts'))[:100] or None,
            'category': _text(data.get('category'))[:200] or None,
        }
        if values['email'] and '@' not in values['email']:
            errors.append((row, f'«Email»: некорректный адрес «{values["email"]}»'))
        result.errors.extend(errors)
        prepared.append((identifier or None, values))

    if not result.ok:
        return result

    with transaction.atomic():
        for identifier, values in prepared:
            existing = None
            if identifier:
                existing = Supplier.objects.filter(identifier=identifier).first()
            if existing is None:
                existing = Supplier.objects.filter(name__iexact=values['name']).first()
            if existing is None:
                Supplier.objects.create(
                    identifier=identifier,
                    status=SupplierStatusEnum.MODERATED.value[0],
                    author=author,
                    **values,
                )
                result.created += 1
            else:
                existing.identifier = identifier or existing.identifier
                for field, value in values.items():
                    if value:
                        setattr(existing, field, value)
                existing.save()
                result.updated += 1
    return result
