"""Synthetic workbook mirroring the mock Avantage file's structure (see
app/services/xlsx_import/parsers/avantage.py docstring) — NOT the real file
(the one inspected 2026-08-27 was itself already fake/mock data, but the
convention in this repo is synthetic fixtures in git regardless, see
maxi_mall_fixture.py). Reproduces the quirks that make this format
different from Maxi Mall:

- One sheet per file, period lives in the A1 title text, not the sheet
  name/tab.
- Paid/remainder columns can be broken formulas (#REF!) — must not crash,
  must not be silently read as a real zero.
- БИН ("Бик" column F) and phone (column G) have no header of their own
  in the charge-types row — only in the row-2 group header, same as the
  real file.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

# (name, brand, unit, zone, bin_value, phone, rent_charged, rent_paid,
#  util_charged, util_paid, debt_charged, debt_paid, other_charged, other_paid)
_ROWS = [
    ("Fully Paid LLP", "BrandA", "B1", "1/A", "111111111111", "77011110000", 100000, 100000, 20000, 20000, 0, 0, 0, 0),
    ("Partial Payer IP", "BrandB", "B2", "1/A", "222222222222", "", 200000, 50000, 30000, 0, 0, 0, 0, 0),
    ("Credit Balance TOO", "BrandC", "B3", "1/B", "333333333333", "77033330000", 0, 0, 0, 0, -50000, 0, 0, 0),
    # Not in the counterparty cache fixture below — stays unmatched (virtual:).
    ("Unknown To 1C IP", "BrandD", "B4", "1/B", "", "", 80000, 0, 0, 0, 0, 0, 0, 0),
    # Name in the file doesn't match anything in the cache by name at all —
    # only its БИН does (cache fullName is deliberately different). Proves
    # BIN-matching actually rescues a case name-matching alone would miss.
    ("Renamed On Paper TOO", "BrandE", "B5", "1/B", "444444444444", "", 90000, 0, 0, 0, 0, 0, 0, 0),
]

_IDENTITY_HEADERS = ["№", "Помещ.", "Арендатор", "Бренд", "Зона"]
_CHARGE_HEADERS = [
    "Аренда\nначисл. ₸", "КУ\nначисл. ₸", "Долг пр.пер.\nначисл. ₸", "Прочее\nначисл. ₸",
    "ИТОГО\nначисл. ₸",
    "Аренда\nоплач. ₸", "КУ\nоплач. ₸", "Долг пр.пер.\nоплач. ₸", "% с обор./\nПрочее ₸",
    "ИТОГО\nоплач. ₸",
    "Аренда\nост. ₸", "КУ\nост. ₸", "Долг\nост. ₸", "ИТОГО\nост. ₸",
    "% сбора", "Статус", "Примечание",
]
# Column layout: identity 1-5, БИН 6, телефон 7 (both row-2-only headers),
# charges 8-24.
_BIN_COL = 6
_PHONE_COL = 7
_CHARGE_START_COL = 8


def build_workbook_bytes(*, broken_paid_formulas: bool = False, note: str = "") -> bytes:
    """broken_paid_formulas=True writes "#REF!" into every paid/remainder
    cell, reproducing the real mock file's state. note, when set, is
    written into a cell further along row 1 (e.g. "А - сентябрь, КУ -
    июль") to exercise the free-text-note warning."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws["A1"] = "ТРЦ «Avantage» — г. Кыргаулды  |  НАЧИСЛЕНИЕ + СВОД ОПЛАТ — АВГУСТ 2026"
    ws.merge_cells("A1:Z1")
    if note:
        ws.cell(row=1, column=28, value=note)

    ws["A2"] = "АРЕНДАТОР"
    ws.merge_cells("A2:E2")
    ws.cell(row=2, column=_BIN_COL, value="Бик")
    ws.cell(row=2, column=_PHONE_COL, value="номер телефона")
    ws.cell(row=2, column=_CHARGE_START_COL, value="◀  НАЧИСЛЕНО  ▶")

    for col, h in enumerate(_IDENTITY_HEADERS, start=1):
        ws.cell(row=3, column=col, value=h)
    # Columns 6-7 (БИН/телефон) deliberately left blank in row 3 — their
    # only label is the row-2 group header above, same as the real file.
    for col, h in enumerate(_CHARGE_HEADERS, start=_CHARGE_START_COL):
        ws.cell(row=3, column=col, value=h)

    row_num = 4
    for name, brand, unit, zone, bin_value, phone, rc, rp, uc, up, dc, dp, oc, op in _ROWS:
        ws.cell(row=row_num, column=1, value=row_num - 3)
        ws.cell(row=row_num, column=2, value=unit)
        ws.cell(row=row_num, column=3, value=name)
        ws.cell(row=row_num, column=4, value=brand)
        ws.cell(row=row_num, column=5, value=zone)
        if bin_value:
            # Excel stores it as a float, same as the real file (e.g.
            # 123456781233.0), not a string.
            ws.cell(row=row_num, column=_BIN_COL, value=float(bin_value))
        if phone:
            ws.cell(row=row_num, column=_PHONE_COL, value=float(phone))

        c = _CHARGE_START_COL
        ws.cell(row=row_num, column=c, value=rc)
        ws.cell(row=row_num, column=c + 1, value=uc)
        ws.cell(row=row_num, column=c + 2, value=dc)
        ws.cell(row=row_num, column=c + 3, value=oc)
        ws.cell(row=row_num, column=c + 4, value=rc + uc + dc + oc)
        if broken_paid_formulas:
            for offset in range(5, 14):
                ws.cell(row=row_num, column=c + offset, value="#REF!")
        else:
            ws.cell(row=row_num, column=c + 5, value=rp)
            ws.cell(row=row_num, column=c + 6, value=up)
            ws.cell(row=row_num, column=c + 7, value=dp)
            ws.cell(row=row_num, column=c + 8, value=op)
            ws.cell(row=row_num, column=c + 9, value=rp + up + dp + op)
            ws.cell(row=row_num, column=c + 10, value=rc - rp)
            ws.cell(row=row_num, column=c + 11, value=uc - up)
            ws.cell(row=row_num, column=c + 12, value=dc - dp)
            ws.cell(row=row_num, column=c + 13, value=(rc - rp) + (uc - up) + (dc - dp))
        row_num += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


COUNTERPARTY_CACHE_FIXTURE = [
    {"id": "cp-fully-paid", "fullName": "Fully Paid LLP"},
    {"id": "cp-partial", "fullName": "Partial Payer IP"},
    {"id": "cp-credit", "fullName": "Credit Balance TOO"},
    # "Unknown To 1C IP" deliberately absent — stays unmatched.
    # Name doesn't match "Renamed On Paper TOO" at all — only reachable via
    # bin_index, proving BIN-matching actually does something.
    {"id": "cp-bin-match", "fullName": "Legally Different Name TOO", "bin": "444444444444"},
]
