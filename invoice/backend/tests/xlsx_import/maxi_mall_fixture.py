"""Synthetic workbook mirroring the real Maxi Mall file's structure (see
app/services/xlsx_import/parsers/maxi_mall.py docstring) — NOT the real
file itself (real tenant financial data doesn't belong in git). Reproduces
the two quirks that made the real parser tricky:

- "Начисление ИЮНЬ" has no "Статус" column (ИЮЛЬ/АВГУСТ do) — column
  resolution must go by header text, not fixed index.
- A merged "ИТОГО ПО ТРЦ" totals row at the bottom, that must be excluded.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

# (name, brand, unit, zone, rent_charged, rent_paid, util_charged, util_paid,
#  debt_charged, debt_paid, other_charged, other_paid)
_ROWS = [
    ("Fully Paid LLP", "BrandA", "B1", "1/A", 100000, 100000, 20000, 20000, 0, 0, 0, 0),
    ("Partial Payer IP", "BrandB", "B2", "1/A", 200000, 50000, 30000, 0, 0, 0, 0, 0),
    # Credit carried from a prior period: negative charge, remainder also
    # negative — must resolve to PAID (nothing owed), not "unpaid" from a
    # naive charged>0-implies-something-owed read.
    ("Credit Balance TOO", "BrandC", "B3", "1/B", 0, 0, 0, 0, -50000, 0, 0, 0),
    # Not in the counterparty cache fixture below — stays unmatched (virtual:).
    ("Unknown To 1C IP", "BrandD", "B4", "1/B", 80000, 0, 0, 0, 0, 0, 0, 0),
]


def _write_charge_sheet(wb: Workbook, sheet_name: str, month_title: str, include_status_col: bool):
    ws = wb.create_sheet(sheet_name)
    ws["A1"] = f'ТРЦ «MAXI MALL» — г. Усть-Каменогорск  |  НАЧИСЛЕНИЕ + СВОД ОПЛАТ — {month_title}'
    ws.merge_cells("A1:U1")

    ws["A2"] = "АРЕНДАТОР"
    ws.merge_cells("A2:E2")
    ws["F2"] = "◀  НАЧИСЛЕНО  ▶"
    ws.merge_cells("F2:I2")
    ws["K2"] = "◀  ОПЛАЧЕНО  ▶"
    ws.merge_cells("K2:N2")

    headers = [
        "№", "Помещ.", "Арендатор", "Бренд", "Зона",
        "Аренда\nначисл. ₸", "КУ\nначисл. ₸", "Долг пр.пер.\nначисл. ₸", "Прочее\nначисл. ₸",
        "ИТОГО\nначисл. ₸",
        "Аренда\nоплач. ₸", "КУ\nоплач. ₸", "Долг пр.пер.\nоплач. ₸", "% с обор./\nПрочее ₸",
        "ИТОГО\nоплач. ₸",
        "Аренда\nост. ₸", "КУ\nост. ₸", "Долг\nост. ₸", "ИТОГО\nост. ₸",
        "% сбора",
    ]
    if include_status_col:
        headers.append("Статус")
    headers.append("Примечание")

    for col, h in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=h)

    totals = {"rent": 0.0, "utilities": 0.0, "debt": 0.0, "other": 0.0}
    row_num = 4
    for name, brand, unit, zone, rc, rp, uc, up, dc, dp, oc, op in _ROWS:
        ws.cell(row=row_num, column=1, value=row_num - 3)
        ws.cell(row=row_num, column=2, value=unit)
        ws.cell(row=row_num, column=3, value=name)
        ws.cell(row=row_num, column=4, value=brand)
        ws.cell(row=row_num, column=5, value=zone)
        ws.cell(row=row_num, column=6, value=rc)
        ws.cell(row=row_num, column=7, value=uc)
        ws.cell(row=row_num, column=8, value=dc)
        ws.cell(row=row_num, column=9, value=oc)
        ws.cell(row=row_num, column=10, value=rc + uc + dc + oc)
        ws.cell(row=row_num, column=11, value=rp)
        ws.cell(row=row_num, column=12, value=up)
        ws.cell(row=row_num, column=13, value=dp)
        ws.cell(row=row_num, column=14, value=op)
        ws.cell(row=row_num, column=15, value=rp + up + dp + op)
        ws.cell(row=row_num, column=16, value=rc - rp)
        ws.cell(row=row_num, column=17, value=uc - up)
        ws.cell(row=row_num, column=18, value=dc - dp)
        ws.cell(row=row_num, column=19, value=(rc - rp) + (uc - up) + (dc - dp))
        totals["rent"] += rc
        totals["utilities"] += uc
        totals["debt"] += dc
        totals["other"] += oc
        row_num += 1

    total_row = row_num
    ws.cell(row=total_row, column=1, value="ИТОГО ПО ТРЦ")
    ws.merge_cells(f"A{total_row}:E{total_row}")
    ws.cell(row=total_row, column=6, value=totals["rent"])
    ws.cell(row=total_row, column=7, value=totals["utilities"])
    ws.cell(row=total_row, column=8, value=totals["debt"])
    ws.cell(row=total_row, column=9, value=totals["other"])

    return totals


def _write_summary_sheet(wb: Workbook, month_label: str, totals: dict):
    ws = wb.create_sheet("Общий свод оплат")
    ws["A4"] = "№"
    ws["B4"] = "Мес."
    ws["C4"] = "Вид платежа"
    ws["D4"] = "Начислено ₸"
    ws["E4"] = "Оплачено ₸"

    rows = [
        (month_label, "Аренда", totals["rent"]),
        (month_label, "КУ", totals["utilities"]),
        (month_label, "Долг пр.пер.", totals["debt"]),
        (month_label, "% с обор./Проч.", totals["other"]),
    ]
    for i, (month, label, charged) in enumerate(rows, start=1):
        r = 4 + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=month)
        ws.cell(row=r, column=3, value=label)
        ws.cell(row=r, column=4, value=charged)


def build_workbook_bytes(*, corrupt_summary: bool = False) -> bytes:
    """Two periods: ИЮНЬ (no Статус column — the real quirk) and ИЮЛЬ (has
    it). corrupt_summary=True deliberately breaks the self-check total to
    exercise the totals_check/warnings path."""
    wb = Workbook()
    wb.remove(wb.active)

    totals_june = _write_charge_sheet(wb, "Начисление ИЮНЬ", "ИЮНЬ 2026", include_status_col=False)
    totals_july = _write_charge_sheet(wb, "Начисление ИЮЛЬ", "ИЮЛЬ 2026", include_status_col=True)

    if corrupt_summary:
        totals_june = dict(totals_june)
        totals_june["rent"] += 999999

    _write_summary_sheet(wb, "ИЮНЬ", totals_june)
    # Реальный файл держит все месяцы в одном "Общий свод оплат" — но по
    # текущей _summary_totals-реализации год/месяц матчатся по метке месяца
    # независимо от листа, так что достаточно дописать вторую пачку строк.
    ws = wb["Общий свод оплат"]
    start_row = ws.max_row + 1
    rows = [
        ("ИЮЛЬ", "Аренда", totals_july["rent"]),
        ("ИЮЛЬ", "КУ", totals_july["utilities"]),
        ("ИЮЛЬ", "Долг пр.пер.", totals_july["debt"]),
        ("ИЮЛЬ", "% с обор./Проч.", totals_july["other"]),
    ]
    for i, (month, label, charged) in enumerate(rows):
        r = start_row + i
        ws.cell(row=r, column=1, value=5 + i)
        ws.cell(row=r, column=2, value=month)
        ws.cell(row=r, column=3, value=label)
        ws.cell(row=r, column=4, value=charged)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


COUNTERPARTY_CACHE_FIXTURE = [
    {"id": "cp-fully-paid", "fullName": "Fully Paid LLP"},
    {"id": "cp-partial", "fullName": "Partial Payer IP"},
    {"id": "cp-credit", "fullName": "Credit Balance TOO"},
    # "Unknown To 1C IP" deliberately absent — stays unmatched.
]
